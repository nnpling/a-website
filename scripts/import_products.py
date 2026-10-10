"""Cập nhật sản phẩm trong src/data/products/ từ file Catsette_Danh_Sach_San_Pham.xlsx.

Cách dùng:
    python3 scripts/import_products.py duong/dan/Catsette_Danh_Sach_San_Pham.xlsx

Script ghi đè tên, danh mục, giá và bảng phân loại theo Excel (kể cả giá đã sửa
trong trang /admin). Tên ngắn, ảnh, link Shopee, thứ tự, sản phẩm nổi bật và tên
nhóm phân loại được giữ nguyên. Sản phẩm không còn trong Excel không bị xoá.
"""
import json
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

OUT = Path(__file__).resolve().parent.parent / "src" / "data" / "products"

# Các ô chỉ sửa trên web, chạy lại script không ghi đè.
KEEP = ("shortName", "order", "featured", "images", "optionImages", "shopeeUrl", "option1Name", "option2Name")

# Sản phẩm đã gỡ khỏi web, chạy lại script không tạo lại.
REMOVED = {
    "the-in-theo-yeu-cau-kem-spotify-hinh-anh-tu-chon",  # trùng card 2 mặt custom
    "khung-anh-nam-cham-trong-suot-dung-polaroid-size",  # ngừng bán
    "card-lyric-mini-2-mat-kem-ma-spotify-taylor-swift",  # ngừng bán
}

CATEGORY_SLUGS = {
    "NFC Keychain": "moc-khoa-nfc",
    "NFC Card": "card-nfc",
    "Photo Card": "photo-card",
    "Lyric Card": "lyric-card",
    "Bookmark": "bookmark",
    "Postcard": "postcard",
    "Poster": "poster",
    "Sticker/Decal": "sticker",
    "Phụ kiện": "phu-kien",
}


def slugify(text: str) -> str:
    text = text.replace("đ", "d").replace("Đ", "D")
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    text = re.sub(r"[^a-zA-Z0-9]+", "-", text).strip("-").lower()
    return text


def parse_price(value):
    """Trả về (min, max) theo đồng; nhận số hoặc chuỗi kiểu '17.000 – 23.000 đ'."""
    if isinstance(value, (int, float)):
        return int(value), int(value)
    nums = [int(n.replace(".", "")) for n in re.findall(r"\d[\d.]*", str(value or ""))]
    return (min(nums), max(nums)) if nums else (None, None)


def split_options(label):
    """Shopee nối 2 phân loại bằng dấu phẩy liền chữ ("Khung tròn,+ móc sao, in 2 mặt")."""
    return [o.strip() for o in re.split(r",(?=\S)", label)][:2]


def unique_slug(name, used):
    base = slugify(re.sub(r"^\[[^\]]*\]\s*", "", name))
    if len(base) > 50:
        base = base[:50].rsplit("-", 1)[0]
    slug, i = base, 2
    while slug in used:
        slug, i = f"{base}-{i}", i + 1
    used.add(slug)
    return slug


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    products, used = [], set()

    ready, custom = wb.worksheets[0], wb.worksheets[1]

    for cat, name, count, price, note in ready.iter_rows(min_row=5, values_only=True):
        if not name:
            continue
        lo, hi = parse_price(price)
        products.append({
            "slug": unique_slug(name, used),
            "name": name.strip(),
            "category": cat,
            "categorySlug": CATEGORY_SLUGS.get(cat, slugify(cat)),
            "type": "ready",
            "variantCount": count,
            "priceMin": lo,
            "priceMax": hi,
            "note": note,
            "variants": [],
        })

    current = None
    for cat, name, variant, price in custom.iter_rows(min_row=5, values_only=True):
        if name:
            current = {
                "slug": unique_slug(name, used),
                "name": name.strip(),
                "category": cat,
                "categorySlug": CATEGORY_SLUGS.get(cat, slugify(cat)),
                "type": "custom",
                "variants": [],
            }
            products.append(current)
        if current is not None and variant:
            current["variants"].append({"label": str(variant).strip(), "price": parse_price(price)[0]})

    for p in products:
        if p["type"] == "custom":
            prices = [v["price"] for v in p["variants"] if v["price"]]
            p["variantCount"] = len(p["variants"])
            p["priceMin"], p["priceMax"] = min(prices), max(prices)
            p["note"] = None

    OUT.mkdir(parents=True, exist_ok=True)
    products = [p for p in products if p["slug"] not in REMOVED]
    for i, p in enumerate(products, 1):
        path = OUT / f"{p['slug']}.json"
        old = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}
        data = {
            "name": p["name"],
            "shortName": p["name"],
            "categorySlug": p["categorySlug"],
            "type": p["type"],
            "order": i,
            "images": [],
            "shopeeUrl": "",
            "note": p["note"] or "",
            "option1Name": "",
            "option2Name": "",
        }
        data.update({k: old[k] for k in KEEP if k in old})
        if old.get("image") and not data["images"]:
            data["images"] = [old["image"]]
        if p["type"] == "ready":
            data.update(priceMin=p["priceMin"], priceMax=p["priceMax"], variantCount=p["variantCount"])
        data["variants"] = [
            dict(zip(("option1", "option2"), split_options(v["label"]) + [""]), price=v["price"])
            for v in p["variants"]
        ]
        path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Đã cập nhật {len(products)} sản phẩm trong {OUT}")


if __name__ == "__main__":
    main(sys.argv[1])
